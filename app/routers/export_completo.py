"""Extração completa do banco — todo mundo que precisa de uma cópia offline dos
dados (backup, análise, levar pra reunião de fora) baixa um único .xlsx com uma
aba por tipo de registro, em vez de garimpar informação tela por tela."""

import io
from datetime import datetime

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, selectinload

from .. import models
from ..database import get_db

router = APIRouter(tags=["export"])

HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")


def _add_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], widths: list[int] | None = None):
    ws = wb.create_sheet(title=title[:31])
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
        r = ws.max_row
        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=r, column=col_idx).alignment = Alignment(vertical="top", wrap_text=True)
    if not rows:
        ws.append(["(nenhum registro)"])
    for col_idx, w in enumerate((widths or [20] * len(headers)), start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w
    ws.freeze_panes = "A2"
    return ws


def _dt(v):
    return v.strftime("%d/%m/%Y %H:%M") if v else ""


def _d(v):
    return v.strftime("%d/%m/%Y") if v else ""


def _bool(v):
    return "Sim" if v else "Não"


@router.get("/export/completo")
def export_completo(db: Session = Depends(get_db)):
    wb = Workbook()
    wb.remove(wb.active)

    # ---------------- capa ----------------
    counts = {
        "Casos de teste": db.query(models.TestCase).count(),
        "Situações": db.query(models.Situacao).count(),
        "Ajustes de ativos": db.query(models.AtivoAjuste).count(),
        "Técnicos (Track One)": db.query(models.Tecnico).count(),
        "Fases do piloto": db.query(models.PilotoFase).count(),
        "Notas de reunião": db.query(models.MeetingNote).count(),
        "Eventos na agenda": db.query(models.AgendaEvento).count(),
        "Tarefas (quadro)": db.query(models.TodoTarefa).count(),
        "Diagramas de fluxo": db.query(models.FlowDiagram).count(),
        "Eventos na trilha de atividades": db.query(models.ActivityLog).count(),
    }
    ws = wb.create_sheet("Resumo")
    ws.append(["Extração completa — Fluxo C (Console de Teste Faiston)"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
    ws.append([])
    ws.append(["Aba", "Total de registros"])
    for col_idx in (1, 2):
        cell = ws.cell(row=4, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    for nome, total in counts.items():
        ws.append([nome, total])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18

    # ---------------- casos de teste ----------------
    cases = (
        db.query(models.TestCase)
        .options(selectinload(models.TestCase.observations), selectinload(models.TestCase.screenshots))
        .order_by(models.TestCase.fluxo, models.TestCase.grupo, models.TestCase.estagio_num.nulls_last(), models.TestCase.code)
        .all()
    )
    _add_sheet(
        wb, "Casos de teste",
        ["Fluxo", "Código", "Grupo", "Estágio", "Nº", "Frente", "Tipo", "Prioridade", "Origem",
         "Status", "Testado por", "Chamado", "Pré-condição", "Passos", "Resultado esperado",
         "Problema encontrado", "Qtd. observações", "Qtd. prints", "Ativo", "Atualizado em"],
        [
            [c.fluxo, c.code, c.grupo, c.estagio, c.estagio_num, c.frente, c.tipo, c.prioridade, c.origem,
             c.status, c.testado_por, c.chamado, c.pre_condicao, c.passos, c.resultado_esperado,
             c.problema_encontrado, len(c.observations), len(c.screenshots), _bool(c.active), _dt(c.updated_at)]
            for c in cases
        ],
        [8, 10, 12, 26, 5, 16, 10, 10, 16, 12, 14, 12, 30, 34, 34, 30, 8, 8, 8, 16],
    )

    obs_rows = []
    for c in cases:
        for o in c.observations:
            obs_rows.append([c.code, o.autor, o.texto, o.cor or "", o.editado_por, _dt(o.editado_em), _dt(o.created_at)])
    _add_sheet(
        wb, "Observações (casos)",
        ["Caso", "Autor", "Texto", "Cor", "Editado por", "Editado em", "Criado em"],
        obs_rows, [10, 16, 46, 10, 16, 16, 16],
    )

    # ---------------- situações ----------------
    situacoes = (
        db.query(models.Situacao)
        .options(selectinload(models.Situacao.estagios).selectinload(models.SituacaoEstagio.observations))
        .order_by(models.Situacao.fluxo, models.Situacao.code)
        .all()
    )
    _add_sheet(
        wb, "Situações",
        ["Fluxo", "Código", "Título", "Descrição", "Origem", "Chamado", "Ativo", "Atualizado em"],
        [
            [s.fluxo, s.code, s.titulo, s.descricao, s.origem, s.chamado, _bool(s.active), _dt(s.updated_at)]
            for s in situacoes
        ],
        [8, 10, 28, 40, 20, 14, 8, 16],
    )

    estagio_rows, sit_obs_rows = [], []
    for s in situacoes:
        for e in s.estagios:
            estagio_rows.append([s.code, e.ordem, e.nome, e.frente, e.passos, e.resultado_esperado,
                                  e.status, e.testado_por, _dt(e.updated_at)])
            for o in e.observations:
                sit_obs_rows.append([s.code, e.nome, o.autor, o.texto, o.cor or "", o.editado_por,
                                      _dt(o.editado_em), _dt(o.created_at)])
    _add_sheet(
        wb, "Estágios das situações",
        ["Situação", "Ordem", "Nome", "Frente", "Passos", "Resultado esperado", "Status", "Testado por", "Atualizado em"],
        estagio_rows, [10, 8, 24, 16, 34, 34, 12, 14, 16],
    )
    _add_sheet(
        wb, "Observações (situações)",
        ["Situação", "Estágio", "Autor", "Texto", "Cor", "Editado por", "Editado em", "Criado em"],
        sit_obs_rows, [10, 24, 16, 46, 10, 16, 16, 16],
    )

    # ---------------- ajustes de ativos ----------------
    ajustes = (
        db.query(models.AtivoAjuste)
        .options(selectinload(models.AtivoAjuste.observations), selectinload(models.AtivoAjuste.prints))
        .order_by(models.AtivoAjuste.versao, models.AtivoAjuste.numero)
        .all()
    )
    _add_sheet(
        wb, "Ajustes de ativos",
        ["Versão", "Nº", "Título", "Tipo", "Área", "Prioridade", "Status", "Responsável", "Autor",
         "Hoje é assim", "Deveria ser assim", "Observação", "Retorno do dev", "Prazo",
         "Qtd. observações", "Qtd. prints", "Atualizado em"],
        [
            [a.versao, a.numero, a.titulo, a.tipo, a.area, a.prioridade, a.status, a.responsavel, a.autor,
             a.atual, a.esperado, a.observacao, a.retorno, a.prazo,
             len(a.observations), len(a.prints), _dt(a.updated_at)]
            for a in ajustes
        ],
        [8, 5, 26, 10, 14, 10, 12, 14, 14, 30, 30, 26, 26, 12, 8, 8, 16],
    )

    aj_obs_rows = []
    for a in ajustes:
        ref = f"{a.versao} #{a.numero:02d}"
        for o in a.observations:
            aj_obs_rows.append([ref, a.titulo, o.autor, o.texto, o.cor or "", o.editado_por, _dt(o.created_at)])
    _add_sheet(
        wb, "Observações (ajustes)",
        ["Ajuste", "Título", "Autor", "Texto", "Cor", "Editado por", "Criado em"],
        aj_obs_rows, [10, 24, 16, 46, 10, 16, 16],
    )

    # ---------------- track one: piloto e técnicos ----------------
    fases = db.query(models.PilotoFase).order_by(models.PilotoFase.ordem).all()
    _add_sheet(
        wb, "Fases do piloto",
        ["Nome", "Descrição", "Status", "Versão do app", "Meta concluídos", "Meta nota", "Meta etapa",
         "Iniciada em", "Liberada em", "Autor"],
        [
            [f.nome, f.descricao, f.status, f.versao_app, f.meta_concluidos, f.meta_nota, f.meta_etapa,
             _dt(f.iniciada_em), _dt(f.liberada_em), f.autor]
            for f in fases
        ],
        [22, 34, 14, 14, 14, 10, 10, 16, 16, 16],
    )

    tecnicos = (
        db.query(models.Tecnico)
        .options(selectinload(models.Tecnico.observacoes), selectinload(models.Tecnico.fase))
        .order_by(models.Tecnico.nome)
        .all()
    )
    _add_sheet(
        wb, "Técnicos",
        ["Nome", "Telefone", "Papel", "Regional", "Líder", "Status", "Fase do piloto", "Nota",
         "Etapas testadas", "Convidado em", "Instalado em", "Concluído em", "Respondido em"],
        [
            [t.nome, t.telefone, t.papel, t.regional, t.lider_nome, t.status,
             t.fase.nome if t.fase else "", t.nota, (t.etapas_testadas or "").replace("|", ", "),
             _dt(t.convidado_em), _dt(t.instalado_em), _dt(t.concluido_em), _dt(t.respondido_em)]
            for t in tecnicos
        ],
        [22, 16, 10, 14, 18, 14, 18, 6, 26, 16, 16, 16, 16],
    )

    tec_obs_rows = []
    for t in tecnicos:
        for o in t.observacoes:
            tec_obs_rows.append([t.nome, o.autor, o.texto, o.tipo or "", o.ajuste_ref or "",
                                  o.versao_app, o.chamado, _dt(o.created_at)])
    _add_sheet(
        wb, "Observações (técnicos)",
        ["Técnico", "Autor", "Texto", "Tipo", "Ajuste gerado", "Versão do app", "Chamado", "Criado em"],
        tec_obs_rows, [22, 16, 44, 12, 14, 14, 14, 16],
    )

    # ---------------- notas, agenda, tarefas, diagramas ----------------
    notas = db.query(models.MeetingNote).order_by(models.MeetingNote.created_at).all()
    _add_sheet(
        wb, "Notas de reunião",
        ["Fluxo", "Estágio", "Texto", "Autor", "Cobrado", "Cobrado em", "Resolvido", "Resolvido em",
         "Retorno", "Prazo"],
        [
            [n.fluxo, n.estagio, n.texto, n.autor, _bool(n.cobrado), _dt(n.cobrado_em),
             _bool(n.resolvido), _dt(n.resolvido_em), n.retorno, n.prazo]
            for n in notas
        ],
        [8, 18, 44, 16, 10, 16, 10, 16, 30, 12],
    )

    eventos = db.query(models.AgendaEvento).order_by(models.AgendaEvento.data).all()
    _add_sheet(
        wb, "Agenda",
        ["Título", "Descrição", "Data", "Início", "Fim", "Tipo", "Concluído", "Autor"],
        [
            [e.titulo, e.descricao, e.data, e.hora_inicio, e.hora_fim, e.tipo, _bool(e.concluido), e.autor]
            for e in eventos
        ],
        [26, 34, 12, 8, 8, 14, 10, 16],
    )

    tarefas = db.query(models.TodoTarefa).order_by(models.TodoTarefa.status, models.TodoTarefa.posicao).all()
    _add_sheet(
        wb, "Tarefas (quadro)",
        ["Título", "Descrição", "Status", "Responsável", "Autor", "Criado em", "Atualizado em"],
        [
            [t.titulo, t.descricao, t.status, t.responsavel, t.autor, _dt(t.created_at), _dt(t.updated_at)]
            for t in tarefas
        ],
        [26, 34, 12, 16, 16, 16, 16],
    )

    diagramas = db.query(models.FlowDiagram).order_by(models.FlowDiagram.fluxo, models.FlowDiagram.ordem).all()
    _add_sheet(
        wb, "Diagramas de fluxo",
        ["Fluxo", "Tipo", "Título", "Descrição", "Código Mermaid", "Atualizado por", "Atualizado em"],
        [
            [d.fluxo, d.kind, d.titulo, d.descricao, d.mermaid, d.atualizado_por, _dt(d.updated_at)]
            for d in diagramas
        ],
        [8, 10, 22, 34, 50, 16, 16],
    )

    # ---------------- trilha de atividades ----------------
    atividades = db.query(models.ActivityLog).order_by(models.ActivityLog.created_at.desc()).all()
    _add_sheet(
        wb, "Trilha de atividades",
        ["Fluxo", "Tipo", "Texto", "Autor", "Caso relacionado", "Criado em"],
        [[a.fluxo, a.tipo, a.texto, a.autor, a.case_code, _dt(a.created_at)] for a in atividades],
        [8, 12, 50, 16, 16, 16],
    )

    wb.move_sheet("Resumo", offset=-len(wb.sheetnames))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"FluxoC_extracao_completa_{datetime.now().strftime('%Y-%m-%d_%H%M')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
