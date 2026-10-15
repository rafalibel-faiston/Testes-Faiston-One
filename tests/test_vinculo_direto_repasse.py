"""Vínculo do técnico na importação: direto, repasse ou CLT.

O que estes testes protegem: todo técnico entra com o código dele; o repasse
entra ligado ao direto (achado pelo código do direto e, sem ele, pelo nome) e,
sem número próprio, com o número do direto — sem que isso vire "telefone
repetido". Repasse com direto que não dá pra achar entra, mas volta no aviso.
"""
from tests.conftest import planilha

CAB = ["Nome", "Telefone", "Tipo", "Codigo", "Direto", "Codigo do Direto"]


def _importa(client, linhas, cabecalho=CAB, **form):
    arquivo = planilha(linhas, cabecalho=cabecalho)
    resp = client.post("/api/tecnicos/importar", files={"file": ("base.xlsx", arquivo)}, data=form)
    assert resp.status_code == 200, resp.text
    return resp.json()


def _por_nome(client):
    return {t["nome"]: t for t in client.get("/api/tecnicos").json()}


def test_direto_repasse_e_clt_com_codigo(client):
    dados = _importa(client, [
        ["Pedro Repasse", "11977776666", "REPASSE", "R-01", "João Direto", "PARC 1020"],  # antes do direto
        ["João Direto", "11999998888", "DIRETO", "PARC 1020", "", ""],
        ["Ana CLT", "11955554444", "CLT", "CLT 7", "", ""],
    ])
    assert dados["rejeitados"] == 0
    assert dados["criados"] == 3
    assert dados["vinculos"] == {
        "direto": 1, "repasse": 1, "clt": 1, "sem_tipo": 0, "sem_codigo": 0, "repasse_sem_direto": 0,
    }
    assert dados["avisos"] == []

    tecs = _por_nome(client)
    direto, repasse, clt = tecs["João Direto"], tecs["Pedro Repasse"], tecs["Ana CLT"]
    assert direto["vinculo"] == "direto" and direto["codigo"] == "PARC 1020"
    assert [r["nome"] for r in direto["repasses"]] == ["Pedro Repasse"]
    assert repasse["vinculo"] == "repasse" and repasse["codigo"] == "R-01"
    assert repasse["direto_id"] == direto["id"]
    assert repasse["direto_nome"] == "João Direto"
    assert repasse["direto_codigo"] == "PARC 1020"
    assert repasse["numero_do_direto"] is False
    assert clt["vinculo"] == "clt" and clt["codigo"] == "CLT 7"
    assert clt["direto_id"] is None and clt["direto_nome"] is None
    # nada de pai/filho na API
    assert "pai_id" not in repasse and "filhos" not in direto


def test_direto_achado_pelo_codigo_mesmo_com_nome_diferente(client):
    _importa(client, [
        ["João da Silva", "11999998888", "DIRETO", "PARC 1020", "", ""],
        ["Pedro", "11977776666", "REPASSE", "R-01", "JOAO SILVA", "PARC 1020"],
    ])
    tecs = _por_nome(client)
    assert tecs["Pedro"]["direto_id"] == tecs["João da Silva"]["id"]
    assert tecs["Pedro"]["direto_nome"] == "JOAO SILVA"   # fica como veio na planilha


def test_direto_achado_pelo_nome_quando_falta_o_codigo(client):
    _importa(client, [
        ["João Direto", "11999998888", "DIRETO", "PARC 1020", "", ""],
        ["Pedro", "11977776666", "REPASSE", "R-01", "joao direto", ""],
    ])
    tecs = _por_nome(client)
    assert tecs["Pedro"]["direto_id"] == tecs["João Direto"]["id"]
    # código do direto completado pelo cadastro dele
    assert tecs["Pedro"]["direto_codigo"] == "PARC 1020"


def test_repasse_sem_numero_entra_com_o_numero_do_direto(client):
    dados = _importa(client, [
        ["João Direto", "11999998888", "DIRETO", "PARC 1020", "", ""],
        ["Pedro", "", "REPASSE", "R-01", "João Direto", "PARC 1020"],
        ["Lucas", "11999998888", "REPASSE", "R-02", "João Direto", "PARC 1020"],  # repetiu o do direto
    ])
    assert dados["rejeitados"] == 0
    assert dados["criados"] == 3
    tecs = _por_nome(client)
    for nome in ("Pedro", "Lucas"):
        assert tecs[nome]["telefone"] == tecs["João Direto"]["telefone"]
        assert tecs[nome]["numero_do_direto"] is True


def test_repasse_sem_numero_usa_celular_do_direto(client):
    cab = CAB + ["Celular do Direto"]
    dados = _importa(client, [
        ["Pedro", "", "REPASSE", "R-01", "Fulano Fora", "PARC 9", "(11) 93333-2222"],
    ], cabecalho=cab)
    assert dados["rejeitados"] == 0
    tec = _por_nome(client)["Pedro"]
    assert tec["telefone"] == "5511933332222"
    assert tec["direto_id"] is None
    assert tec["direto_nome"] == "Fulano Fora"
    assert tec["direto_codigo"] == "PARC 9"


def test_repasse_sem_numero_e_sem_direto_achado_e_rejeitado(client):
    dados = _importa(client, [["Pedro", "", "REPASSE", "R-01", "Ninguém", "X"]])
    assert dados["rejeitados"] == 1
    assert "sem telefone" in dados["erros"][0]["motivo"]


def test_repasse_com_direto_fora_da_base_vira_aviso(client):
    dados = _importa(client, [
        ["Pedro", "11977776666", "REPASSE", "R-01", "Fulano Fora", "PARC 9"],
        ["Lucas", "11966665555", "REPASSE", "R-02", "", ""],
    ])
    assert dados["criados"] == 2
    assert dados["vinculos"]["repasse_sem_direto"] == 2
    motivos = {a["nome"]: a["motivo"] for a in dados["avisos"]}
    assert "não encontrado" in motivos["Pedro"] and "PARC 9" in motivos["Pedro"]
    assert motivos["Lucas"] == "repasse sem direto informado"
    tec = _por_nome(client)["Pedro"]
    assert tec["direto_nome"] == "Fulano Fora" and tec["direto_codigo"] == "PARC 9"


def test_repasse_acha_direto_que_ja_estava_na_base(client):
    _importa(client, [["João Direto", "11999998888", "DIRETO", "PARC 1020", "", ""]])
    _importa(client, [["Pedro", "11977776666", "REPASSE", "R-01", "", "PARC 1020"]])
    tecs = _por_nome(client)
    assert tecs["Pedro"]["direto_id"] == tecs["João Direto"]["id"]
    assert tecs["Pedro"]["direto_nome"] == "João Direto"


def test_dois_diretos_com_o_mesmo_numero_continua_rejeitado(client):
    dados = _importa(client, [
        ["Ana", "11988887777", "DIRETO", "PARC 1", "", ""],
        ["Bruno", "11988887777", "DIRETO", "PARC 2", "", ""],
    ])
    assert dados["criados"] == 1
    assert dados["rejeitados"] == 1
    assert "repetido" in dados["erros"][0]["motivo"]


def test_sem_codigo_conta_no_resumo(client):
    dados = _importa(client, [["Ana", "11988887777", "CLT", "", "", ""]])
    assert dados["vinculos"]["sem_codigo"] == 1


def test_reimportar_nao_duplica_e_casa_pelo_codigo(client):
    linhas = [
        ["João Direto", "11999998888", "DIRETO", "PARC 1020", "", ""],
        ["Pedro", "", "REPASSE", "R-01", "João Direto", "PARC 1020"],
    ]
    _importa(client, linhas)
    # 2ª importação: o direto trocou de número — o código é o mesmo
    dados = _importa(client, [
        ["João Direto", "11911112222", "DIRETO", "PARC 1020", "", ""],
        ["Pedro", "", "REPASSE", "R-01", "João Direto", "PARC 1020"],
    ])
    assert dados["criados"] == 0
    assert dados["atualizados"] == 2
    tecs = _por_nome(client)
    assert len(tecs) == 2
    assert tecs["João Direto"]["telefone"] == "5511911112222"
    assert tecs["Pedro"]["telefone"] == "5511911112222"
    assert tecs["Pedro"]["direto_id"] == tecs["João Direto"]["id"]


def test_tipo_aceita_variacoes(client):
    dados = _importa(client, [
        ["A", "11911111111", "Parc Direto", "1", "", ""],
        ["B", "11922222222", "PARC REPASSE", "2", "A", "1"],
        ["C", "11933333333", "clt", "3", "", ""],
    ])
    assert dados["vinculos"]["direto"] == 1
    assert dados["vinculos"]["repasse"] == 1
    assert dados["vinculos"]["clt"] == 1


def test_coluna_celular_do_direto_nao_vira_o_telefone(client):
    # sem coluna "Telefone" exata, o "contém" não pode pegar o celular do direto
    cab = ["Nome", "Celular do Técnico", "Tipo", "Codigo", "Direto", "Celular do Direto"]
    _importa(client, [["Pedro", "11977776666", "REPASSE", "R-01", "Fulano", "11933332222"]], cabecalho=cab)
    assert _por_nome(client)["Pedro"]["telefone"] == "5511977776666"


def _base(client, **params):
    return {t["nome"] for t in client.get("/api/tecnicos/base", params=params).json()["itens"]}


def test_filtro_por_vinculo_na_base_e_na_lista(client):
    _importa(client, [
        ["João Direto", "11999998888", "DIRETO", "PARC 1020", "", ""],
        ["Pedro", "11977776666", "REPASSE", "R-01", "João Direto", "PARC 1020"],
        ["Lucas", "11966665555", "REPASSE", "R-02", "Fora", "X"],
        ["Ana", "11955554444", "CLT", "C-1", "", ""],
    ])
    _importa(client, [["Sem Tipo", "11944443333"]], cabecalho=["Nome", "Telefone"])
    assert _base(client, vinculo="direto") == {"João Direto"}
    assert _base(client, vinculo="repasse") == {"Pedro", "Lucas"}
    assert _base(client, vinculo="repasse_sem_direto") == {"Lucas"}
    assert _base(client, vinculo="clt") == {"Ana"}
    assert _base(client, vinculo="sem_vinculo") == {"Sem Tipo"}
    lista = {t["nome"] for t in client.get("/api/tecnicos", params={"vinculo": "repasse"}).json()}
    assert lista == {"Pedro", "Lucas"}

    item = next(t for t in client.get("/api/tecnicos/base").json()["itens"] if t["nome"] == "João Direto")
    assert item["repasses_qtd"] == 1
    assert client.get("/api/tecnicos/base", params={"vinculo": "pai"}).status_code == 400


def test_busca_pelo_codigo_do_direto_acha_os_repasses(client):
    _importa(client, [
        ["João Direto", "11999998888", "DIRETO", "PARC 1020", "", ""],
        ["Pedro", "11977776666", "REPASSE", "R-01", "João Direto", "PARC 1020"],
        ["Ana", "11955554444", "CLT", "C-1", "", ""],
    ])
    assert _base(client, busca="PARC 1020") == {"João Direto", "Pedro"}
    assert _base(client, busca="R-01") == {"Pedro"}


def test_adicionar_na_fase_todos_os_repasses_do_filtro(client):
    _importa(client, [
        ["João Direto", "11999998888", "DIRETO", "PARC 1020", "", ""],
        ["Pedro", "11977776666", "REPASSE", "R-01", "João Direto", "PARC 1020"],
    ])
    fase = client.post("/api/piloto/fases", json={"nome": "Só repasse"}).json()
    resp = client.post(f"/api/piloto/fases/{fase['id']}/tecnicos", json={"vinculo": "repasse"})
    assert resp.json()["adicionados"] == 1


def test_exportar_traz_vinculo_codigo_e_direto(client):
    from io import BytesIO
    from openpyxl import load_workbook
    _importa(client, [
        ["João Direto", "11999998888", "DIRETO", "PARC 1020", "", ""],
        ["Pedro", "", "REPASSE", "R-01", "João Direto", "PARC 1020"],
    ])
    wb = load_workbook(BytesIO(client.get("/api/tecnicos/exportar").content))
    linhas = list(wb["Técnicos"].iter_rows(values_only=True))
    cab = linhas[0]
    for col in ("Vínculo", "Código", "Direto", "Código do direto"):
        assert col in cab
    pedro = dict(zip(cab, next(l for l in linhas if l[0] == "Pedro")))
    assert pedro["Vínculo"] == "Repasse"
    assert pedro["Código"] == "R-01"
    assert pedro["Direto"] == "João Direto"
    assert pedro["Código do direto"] == "PARC 1020"


def test_criar_e_editar_vinculo_na_mao(client):
    direto = client.post("/api/tecnicos", json={
        "nome": "João Direto", "telefone": "11999998888", "vinculo": "direto", "codigo": "PARC 1020",
    }).json()
    rep = client.post("/api/tecnicos", json={
        "nome": "Pedro", "telefone": "11977776666", "vinculo": "repasse", "codigo": "R-01",
        "direto_codigo": "PARC 1020",
    }).json()
    assert rep["direto_id"] == direto["id"]
    assert rep["direto_nome"] == "João Direto"

    # virou CLT: perde o direto
    clt = client.patch(f"/api/tecnicos/{rep['id']}", json={"vinculo": "clt"}).json()
    assert clt["vinculo"] == "clt"
    assert clt["direto_id"] is None and clt["direto_nome"] is None
    assert client.patch(f"/api/tecnicos/{rep['id']}", json={"vinculo": "pai"}).status_code == 400


CAB_AGENTES = ["Técnico", "Tipo", "UF", "Agente", "Telefone"]


def test_planilha_da_base_de_agentes(client):
    """Formato que sobe de verdade: Técnico | Tipo | UF | Agente | Telefone. O
    PARC do Agente é o código do direto; no repasse, é o PARC do direto dele."""
    dados = _importa(client, [
        ["CIRO MARTINHO ROCHA", "Repasse", "ES", "ES PARC1700 REPASSE", 27995257845],
        ["EDIVALDO DE SOUSA LIMA LEAL", "Direto", "ES", "ES PARC1700", 27988948231],
        ["ELDER TELAROLLI", "Repasse", "ES", "ES  PARC1717 REPASSE", 27997782658],
        ["IGOR SALOMAO", "Repasse", "ES", "ES PARC0856 REPASSE", 13981900175],
        ["MARCOS ANTONIO INTRA", "Repasse", "ES", "ES PARC0856 REPASSE", 13981900175],  # mesmo nº
        ["ADNER DA SILVA", "Repasse", "MG", "MG PARC1717 REPASSE", 3591328323],
        ["ANA CLT", "CLT", "MG", "", 31999998888],
    ], cabecalho=CAB_AGENTES)
    assert dados["rejeitados"] == 0, dados["erros"]
    assert dados["criados"] == 7
    assert dados["vinculos"]["direto"] == 1
    assert dados["vinculos"]["repasse"] == 5
    assert dados["vinculos"]["clt"] == 1
    assert dados["vinculos"]["sem_codigo"] == 1          # só o CLT sem Agente
    # PARC1717 e PARC0856 não estão na planilha nem na base
    assert dados["vinculos"]["repasse_sem_direto"] == 4

    tecs = _por_nome(client)
    direto = tecs["EDIVALDO DE SOUSA LIMA LEAL"]
    assert direto["vinculo"] == "direto" and direto["codigo"] == "PARC1700"
    assert direto["regional"] == "ES"
    assert direto["telefone"] == "5527988948231"
    ciro = tecs["CIRO MARTINHO ROCHA"]
    assert ciro["vinculo"] == "repasse"
    assert ciro["codigo"] == "PARC1700 REPASSE"
    assert ciro["direto_id"] == direto["id"]
    assert ciro["direto_codigo"] == "PARC1700"
    assert ciro["direto_nome"] == "EDIVALDO DE SOUSA LIMA LEAL"
    elder = tecs["ELDER TELAROLLI"]
    assert elder["codigo"] == "PARC1717 REPASSE" and elder["direto_codigo"] == "PARC1717"
    assert tecs["ADNER DA SILVA"]["telefone"] == "553591328323"


def test_agentes_repasse_acha_direto_de_outra_planilha(client):
    # direto subiu na planilha de outra região antes
    _importa(client, [["DIRETO 0856", "Direto", "SP", "SP PARC0856", 13981900175]], cabecalho=CAB_AGENTES)
    _importa(client, [["IGOR SALOMAO", "Repasse", "ES", "ES PARC0856 REPASSE", 13981900175]],
             cabecalho=CAB_AGENTES)
    tecs = _por_nome(client)
    igor = tecs["IGOR SALOMAO"]
    assert igor["direto_id"] == tecs["DIRETO 0856"]["id"]
    assert igor["numero_do_direto"] is True


def test_agentes_reimportar_nao_duplica(client):
    linhas = [
        ["EDIVALDO", "Direto", "ES", "ES PARC1700", 27988948231],
        ["CIRO", "Repasse", "ES", "ES PARC1700 REPASSE", 27995257845],
        ["IGOR", "Repasse", "ES", "ES PARC1700 REPASSE", 27995257845],
    ]
    _importa(client, linhas, cabecalho=CAB_AGENTES)
    dados = _importa(client, linhas, cabecalho=CAB_AGENTES)
    assert dados["criados"] == 0 and dados["atualizados"] == 3
    assert len(_por_nome(client)) == 3


def test_direto_importado_depois_liga_os_repasses_que_ja_estavam(client):
    """Repasse subiu numa planilha e o direto noutra, depois: a ligação se
    acerta no fim da importação do direto."""
    _importa(client, [
        ["IGOR SALOMAO", "Repasse", "ES", "ES PARC0856 REPASSE", "13981900175"],
        ["MARCOS ANTONIO INTRA", "Repasse", "MG", "MG PARC0856 REPASSE", "13981900175"],
    ], cabecalho=CAB_AGENTES)
    assert all(t["direto_id"] is None for t in client.get("/api/tecnicos").json())

    dados = _importa(client, [["VALMIR DIAS RIBEIRO", "Direto", "SP", "SP PARC0856", "13981900175"]],
                     cabecalho=CAB_AGENTES)
    assert dados["religacao"]["religados"] == 2
    tecs = _por_nome(client)
    valmir = tecs["VALMIR DIAS RIBEIRO"]
    assert sorted(r["nome"] for r in valmir["repasses"]) == ["IGOR SALOMAO", "MARCOS ANTONIO INTRA"]
    igor = tecs["IGOR SALOMAO"]
    assert igor["direto_id"] == valmir["id"]
    assert igor["direto_nome"] == "VALMIR DIAS RIBEIRO"
    assert igor["numero_do_direto"] is True


def test_religar_desempata_codigo_repetido_pelo_numero(client):
    _importa(client, [
        ["DIRETO A", "Direto", "SP", "SP PARC0856", "11911111111"],
        ["DIRETO B", "Direto", "MG", "MG PARC0856", "13981900175"],
    ], cabecalho=CAB_AGENTES)
    _importa(client, [["IGOR", "Repasse", "ES", "ES PARC0856 REPASSE", "13981900175"]], cabecalho=CAB_AGENTES)
    tecs = _por_nome(client)
    assert tecs["IGOR"]["direto_id"] == tecs["DIRETO B"]["id"]


def test_religar_pelo_endpoint(client):
    _importa(client, [["IGOR", "Repasse", "ES", "ES PARC0856 REPASSE", "13981900175"]], cabecalho=CAB_AGENTES)
    client.post("/api/tecnicos", json={"nome": "VALMIR", "telefone": "13981900175",
                                       "vinculo": "direto", "codigo": "PARC0856"})
    assert client.post("/api/tecnicos/religar").json() == {"religados": 1, "ambiguos": []}
    assert _por_nome(client)["IGOR"]["direto_nome"] == "VALMIR"


def test_convite_do_direto_pergunta_se_pode_chamar_os_repasses(client):
    _importa(client, [
        ["VALMIR DIAS RIBEIRO", "Direto", "SP", "SP PARC0856", "13981900175"],
        ["IGOR SALOMAO", "Repasse", "ES", "ES PARC0856 REPASSE", "13981900175"],
        ["MARCOS ANTONIO INTRA", "Repasse", "ES", "ES PARC0856 REPASSE", "13981900175"],
        ["ANA SOZINHA", "Direto", "SP", "SP PARC0001", "11911112222"],
    ], cabecalho=CAB_AGENTES)
    tecs = _por_nome(client)

    for tipo in ("convite", "acionamento"):
        msg = client.get(f"/api/tecnicos/{tecs['VALMIR DIAS RIBEIRO']['id']}/mensagem",
                         params={"tipo": tipo}).json()["mensagem"]
        assert "técnicos de repasse trabalhando com você: Igor Salomao, Marcos Antonio Intra." in msg
        assert "Podemos mandar esse mesmo convite pra eles também?" in msg
        # entra antes do fecho da mensagem
        assert msg.index("Podemos mandar") < msg.index("Qualquer dúvida")

    # direto sem repasse e o próprio repasse recebem o convite normal
    for nome in ("ANA SOZINHA", "IGOR SALOMAO"):
        msg = client.get(f"/api/tecnicos/{tecs[nome]['id']}/mensagem").json()["mensagem"]
        assert "técnicos de repasse" not in msg


def test_convite_do_direto_com_muitos_repasses_resume(client):
    linhas = [["DIRETO GRANDE", "Direto", "SP", "SP PARC0900", "11900000000"]]
    linhas += [[f"REPASSE {i:02d}", "Repasse", "SP", "SP PARC0900 REPASSE", f"119100000{i:02d}"] for i in range(10)]
    _importa(client, linhas, cabecalho=CAB_AGENTES)
    direto = _por_nome(client)["DIRETO GRANDE"]
    msg = client.get(f"/api/tecnicos/{direto['id']}/mensagem").json()["mensagem"]
    assert "Repasse 07 e mais 2." in msg


def test_convite_fala_da_regiao_do_tecnico(client):
    _importa(client, [
        ["ANA SP", "Direto", "SP", "SP PARC1", "11911111111"],
        ["BIA MG", "Direto", "MG", "MG PARC2", "31922222222"],
        ["CAIO ES", "Direto", "ES", "ES PARC3", "27933333333"],
        ["DUDA RJ", "Direto", "RJ", "RJ PARC4", "21944444444"],
    ], cabecalho=CAB_AGENTES)
    client.post("/api/tecnicos", json={"nome": "EDU SEM UF", "telefone": "71955555555", "regional": "Sudeste"})
    esperado = {
        "ANA SP": "chamados reais da Grande São Paulo passarão",
        "BIA MG": "chamados reais de Minas Gerais passarão",
        "CAIO ES": "chamados reais do Espírito Santo passarão",
        "DUDA RJ": "chamados reais do Rio de Janeiro passarão",
        "EDU SEM UF": "chamados reais da sua região passarão",
    }
    for nome, trecho in esperado.items():
        tid = _por_nome(client)[nome]["id"]
        for tipo in ("convite", "acionamento"):
            msg = client.get(f"/api/tecnicos/{tid}/mensagem", params={"tipo": tipo}).json()["mensagem"]
            assert trecho in msg, (nome, tipo)
