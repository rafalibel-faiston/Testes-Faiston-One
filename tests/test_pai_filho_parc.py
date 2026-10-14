"""Hierarquia pai (PARC) / filho (PARC REPASSE) na importação da base.

O que estes testes protegem: na base de Agentes o filho de repasse costuma vir
com o número do pai. Número repetido + REPASSE = filho, e o número não é dele;
número repetido + só PARC = o dono verdadeiro do número, o pai — que precisa
listar todos os filhos dele.
"""
from tests.conftest import planilha

CAB = ["Nome", "Numero", "Codigo de PARC"]


def _importa(client, linhas, **form):
    arquivo = planilha(linhas, cabecalho=CAB)
    resp = client.post("/api/tecnicos/importar", files={"file": ("base.xlsx", arquivo)}, data=form)
    assert resp.status_code == 200, resp.text
    return resp.json()


def _por_nome(client):
    return {t["nome"]: t for t in client.get("/api/tecnicos").json()}


def test_numero_repetido_com_repasse_vira_filho_do_parc(client):
    dados = _importa(client, [
        ["Pedro Filho", "11999998888", "PARC 1020 REPASSE"],   # filho veio antes do pai
        ["João Pai", "11999998888", "PARC 1020"],
        ["Lucas Filho", "11999998888", "PARC 1020 REPASSE"],
        ["Marcos Solo", "11977776666", "PARC 3000"],
    ])
    assert dados["rejeitados"] == 0
    assert dados["criados"] == 4
    assert dados["hierarquia"] == {
        "pais_com_filhos": 1, "filhos": 2, "filhos_com_numero_do_pai": 2, "filhos_sem_pai": 0,
    }

    tecs = _por_nome(client)
    pai = tecs["João Pai"]
    assert pai["repasse"] is False
    assert pai["numero_do_pai"] is False
    assert pai["pai_id"] is None
    assert sorted(f["nome"] for f in pai["filhos"]) == ["Lucas Filho", "Pedro Filho"]

    for nome in ("Pedro Filho", "Lucas Filho"):
        filho = tecs[nome]
        assert filho["repasse"] is True
        assert filho["pai_id"] == pai["id"]
        assert filho["pai_nome"] == "João Pai"
        # o número gravado é o do pai, sinalizado como não sendo dele
        assert filho["numero_do_pai"] is True
        assert filho["telefone"] == pai["telefone"]
        assert filho["codigo_parc"] == "PARC 1020 REPASSE"

    assert tecs["Marcos Solo"]["filhos"] == []


def test_dois_parc_com_o_mesmo_numero_continua_rejeitado(client):
    dados = _importa(client, [
        ["Ana", "11988887777", "PARC 1"],
        ["Bruno", "11988887777", "PARC 2"],
    ])
    assert dados["criados"] == 1
    assert dados["resumo"] == {"telefone repetido na planilha": 1}
    assert dados["erros"][0]["linha"] == 3


def test_filho_com_numero_proprio_fica_com_o_numero(client):
    dados = _importa(client, [["Rita", "11955554444", "PARC 9 REPASSE"]])
    assert dados["hierarquia"]["filhos_sem_pai"] == 0
    rita = _por_nome(client)["Rita"]
    assert rita["repasse"] is True
    assert rita["numero_do_pai"] is False
    assert rita["pai_id"] is None


def test_filhos_repetidos_sem_parc_pai_viram_aviso(client):
    dados = _importa(client, [
        ["Filho A", "11944443333", "PARC 5 REPASSE"],
        ["Filho B", "11944443333", "PARC 5 REPASSE"],
    ])
    assert dados["criados"] == 2
    assert dados["hierarquia"]["filhos_sem_pai"] == 2
    assert {a["linha"] for a in dados["avisos"]} == {2, 3}
    tecs = _por_nome(client)
    assert tecs["Filho A"]["numero_do_pai"] is True
    assert tecs["Filho A"]["pai_id"] is None


def test_filho_acha_pai_que_ja_estava_na_base(client):
    _importa(client, [["João Pai", "11999998888", "PARC 1020"]])
    dados = _importa(client, [["Pedro Filho", "11999998888", "PARC 1020 REPASSE"]])
    assert dados["criados"] == 1
    tecs = _por_nome(client)
    assert tecs["Pedro Filho"]["pai_id"] == tecs["João Pai"]["id"]
    assert tecs["Pedro Filho"]["numero_do_pai"] is True


def test_reimportar_nao_duplica_pai_nem_filhos(client):
    linhas = [
        ["João Pai", "11999998888", "PARC 1020"],
        ["Pedro Filho", "11999998888", "PARC 1020 REPASSE"],
    ]
    _importa(client, linhas)
    dados = _importa(client, linhas)
    assert dados["criados"] == 0
    assert dados["atualizados"] == 2
    assert len(client.get("/api/tecnicos").json()) == 2


def test_pai_e_filhos_entram_na_fase(client):
    fase = client.post("/api/piloto/fases", json={"nome": "Fase 3"}).json()
    dados = _importa(client, [
        ["João Pai", "11999998888", "PARC 1020"],
        ["Pedro Filho", "11999998888", "PARC 1020 REPASSE"],
    ], fase_id=str(fase["id"]))
    assert dados["fases"][0]["tecnicos"] == 2
    na_fase = client.get(f"/api/tecnicos?fase_id={fase['id']}").json()
    assert {t["nome"] for t in na_fase} == {"João Pai", "Pedro Filho"}


def test_exportar_mostra_vinculo_e_tira_numero_do_filho(client):
    import io

    from openpyxl import load_workbook

    _importa(client, [
        ["João Pai", "11999998888", "PARC 1020"],
        ["Pedro Filho", "11999998888", "PARC 1020 REPASSE"],
    ])
    resp = client.get("/api/tecnicos/exportar")
    ws = load_workbook(io.BytesIO(resp.content)).active
    linhas = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert linhas["Pedro Filho"][1] in (None, "")          # sem número próprio
    assert linhas["Pedro Filho"][11] == "Filho de João Pai"
    assert linhas["Pedro Filho"][12] == "5511999998888"      # só o número do pai
    assert linhas["João Pai"][11] == "Pai de 1: Pedro Filho"


def test_lista_da_base_mostra_quem_e_pai_e_quem_e_filho(client):
    _importa(client, [
        ["João Pai", "11999998888", "PARC 1020"],
        ["Pedro Filho", "11999998888", "PARC 1020 REPASSE"],
    ])
    itens = {t["nome"]: t for t in client.get("/api/tecnicos/base").json()["itens"]}
    assert itens["João Pai"]["filhos_qtd"] == 1
    assert itens["Pedro Filho"]["pai_nome"] == "João Pai"
    assert itens["Pedro Filho"]["numero_do_pai"] is True


def _base_vinculo(client):
    _importa(client, [
        ["João Pai", "11999998888", "PARC 1020"],
        ["Pedro Filho", "11999998888", "PARC 1020 REPASSE"],
        ["Órfão A", "11944443333", "PARC 5 REPASSE"],
        ["Órfão B", "11944443333", "PARC 5 REPASSE"],
        ["Marcos Solo", "11977776666", "PARC 3000"],
    ])


def test_filtro_por_vinculo_na_base_e_na_lista(client):
    _base_vinculo(client)
    esperado = {
        "pai": {"João Pai"},
        "filho": {"Pedro Filho", "Órfão A", "Órfão B"},
        "filho_sem_pai": {"Órfão A", "Órfão B"},
        "sem_vinculo": {"Marcos Solo"},
    }
    for vinculo, nomes in esperado.items():
        base = client.get(f"/api/tecnicos/base?vinculo={vinculo}").json()
        assert {t["nome"] for t in base["itens"]} == nomes, vinculo
        assert base["total"] == len(nomes)
        lista = client.get(f"/api/tecnicos?vinculo={vinculo}").json()
        assert {t["nome"] for t in lista} == nomes, vinculo


def test_filtro_de_vinculo_invalido_da_400(client):
    assert client.get("/api/tecnicos/base?vinculo=tio").status_code == 400


def test_busca_acha_pelo_codigo_parc(client):
    _base_vinculo(client)
    itens = client.get("/api/tecnicos/base?busca=PARC 1020").json()["itens"]
    assert {t["nome"] for t in itens} == {"João Pai", "Pedro Filho"}


def test_adicionar_na_fase_todos_os_filhos_do_filtro(client):
    _base_vinculo(client)
    fase = client.post("/api/piloto/fases", json={"nome": "Só filhos"}).json()
    resp = client.post(f"/api/piloto/fases/{fase['id']}/tecnicos", json={"vinculo": "filho"})
    assert resp.json()["adicionados"] == 3
    na_fase = client.get(f"/api/tecnicos?fase_id={fase['id']}").json()
    assert {t["nome"] for t in na_fase} == {"Pedro Filho", "Órfão A", "Órfão B"}
