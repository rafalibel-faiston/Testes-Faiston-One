"""Divisão da fase entre as pessoas do time, por UF.

O que estes testes protegem: as filas ficam do mesmo tamanho, cada UF fica
inteira com uma pessoa sempre que cabe, a UF grande é repartida mantendo o
direto junto dos repasses dele, e rodar de novo não embaralha quem já tinha
dono (a pessoa já pode ter mandado mensagem pra ele).
"""
from tests.conftest import planilha

CAB = ["Técnico", "Tipo", "UF", "Agente", "Telefone"]
PESSOAS = ["Luis", "Rafael", "Renato"]


def _fase_com(client, por_uf):
    fase = client.post("/api/piloto/fases", json={"nome": "Fase 1"}).json()
    linhas, n = [], 0
    for uf, qtd in por_uf.items():
        for _ in range(qtd):
            n += 1
            linhas.append([f"TEC {uf} {n:03d}", "Direto", uf, f"{uf} PARC{n}", 11900000000 + n])
    client.post("/api/tecnicos/importar", files={"file": ("b.xlsx", planilha(linhas, cabecalho=CAB))},
                data={"fase_id": fase["id"]})
    return fase


def _dividir(client, fase, **extra):
    resp = client.post(f"/api/piloto/fases/{fase['id']}/dividir", json={"pessoas": PESSOAS, **extra})
    assert resp.status_code == 200, resp.text
    return resp.json()


def _filas(client, fase):
    filas = {}
    for t in client.get("/api/tecnicos", params={"fase_id": fase["id"]}).json():
        filas.setdefault(t["responsavel"], []).append(t)
    return filas


def test_divide_igual_e_mantem_uf_inteira_quando_cabe(client):
    fase = _fase_com(client, {"MG": 10, "ES": 10, "RJ": 10})
    dados = _dividir(client, fase)
    assert dados["divididos"] == 30
    filas = _filas(client, fase)
    assert sorted(len(v) for v in filas.values()) == [10, 10, 10]
    for fila in filas.values():
        assert len({t["regional"] for t in fila}) == 1      # uma UF por pessoa


def test_uf_grande_e_repartida_pra_equilibrar(client):
    fase = _fase_com(client, {"SP": 30, "MG": 6, "ES": 3})
    _dividir(client, fase)
    filas = _filas(client, fase)
    assert sorted(len(v) for v in filas.values()) == [13, 13, 13]
    # MG e ES continuam inteiras, cada uma com uma pessoa só
    for uf in ("MG", "ES"):
        assert sum(1 for fila in filas.values() if any(t["regional"] == uf for t in fila)) == 1


def test_direto_e_repasses_ficam_com_a_mesma_pessoa(client):
    fase = client.post("/api/piloto/fases", json={"nome": "F"}).json()
    linhas = []
    for d in range(6):
        linhas.append([f"DIRETO {d}", "Direto", "SP", f"SP PARC{100 + d}", 11910000000 + d])
        for r in range(2):
            linhas.append([f"REP {d}-{r}", "Repasse", "SP", f"SP PARC{100 + d} REPASSE", 11920000000 + d * 10 + r])
    client.post("/api/tecnicos/importar", files={"file": ("b.xlsx", planilha(linhas, cabecalho=CAB))},
                data={"fase_id": fase["id"]})
    _dividir(client, fase)
    tecs = {t["nome"]: t for t in client.get("/api/tecnicos", params={"fase_id": fase["id"]}).json()}
    for d in range(6):
        assert tecs[f"REP {d}-0"]["responsavel"] == tecs[f"DIRETO {d}"]["responsavel"]
        assert tecs[f"REP {d}-1"]["responsavel"] == tecs[f"DIRETO {d}"]["responsavel"]


def test_rodar_de_novo_nao_mexe_em_quem_ja_tem_dono(client):
    fase = _fase_com(client, {"MG": 9})
    _dividir(client, fase)
    antes = {t["id"]: t["responsavel"] for t in client.get("/api/tecnicos", params={"fase_id": fase["id"]}).json()}
    dados = _dividir(client, fase)
    assert dados["divididos"] == 0
    depois = {t["id"]: t["responsavel"] for t in client.get("/api/tecnicos", params={"fase_id": fase["id"]}).json()}
    assert antes == depois


def test_novos_na_fase_vao_pra_quem_tem_menos(client):
    fase = _fase_com(client, {"MG": 6})
    client.post(f"/api/piloto/fases/{fase['id']}/dividir", json={"pessoas": ["Luis", "Rafael"]})
    dados = _dividir(client, fase)          # entrou o Renato
    assert dados["divididos"] == 0
    # ninguém tirado de ninguém: o Renato só pega quando chegar gente sem dono
    novos = [["NOVO 1", "Direto", "MG", "MG PARC900", 11980000001],
             ["NOVO 2", "Direto", "MG", "MG PARC901", 11980000002]]
    client.post("/api/tecnicos/importar", files={"file": ("b.xlsx", planilha(novos, cabecalho=CAB))},
                data={"fase_id": fase["id"]})
    _dividir(client, fase)
    filas = _filas(client, fase)
    assert len(filas["Renato"]) == 2


def test_redividir_tudo_e_grafia_da_lista(client):
    fase = _fase_com(client, {"MG": 6})
    tec = client.get("/api/tecnicos", params={"fase_id": fase["id"]}).json()[0]
    client.patch(f"/api/tecnicos/{tec['id']}", json={"responsavel": "luís"})
    _dividir(client, fase)
    assert client.get(f"/api/tecnicos", params={"fase_id": fase["id"]}).json()[0]["responsavel"] == "Luis"
    dados = _dividir(client, fase, manter_atribuidos=False)
    assert dados["divididos"] == 6
    assert [p["tecnicos"] for p in dados["pessoas"]] == [2, 2, 2]


def test_filtro_por_responsavel_e_export(client):
    from io import BytesIO
    from openpyxl import load_workbook
    fase = _fase_com(client, {"MG": 3, "ES": 3, "RJ": 3})
    _dividir(client, fase)
    meus = client.get("/api/tecnicos", params={"fase_id": fase["id"], "responsavel": "Rafael"}).json()
    assert len(meus) == 3 and {t["responsavel"] for t in meus} == {"Rafael"}
    wb = load_workbook(BytesIO(client.get("/api/tecnicos/exportar").content))
    assert "Responsável" in next(wb["Técnicos"].iter_rows(values_only=True))


def test_sem_pessoas_da_400(client):
    fase = _fase_com(client, {"MG": 1})
    assert client.post(f"/api/piloto/fases/{fase['id']}/dividir", json={"pessoas": [" "]}).status_code == 400
    assert client.post("/api/piloto/fases/99999/dividir", json={"pessoas": PESSOAS}).status_code == 404


def test_mensagem_sai_no_nome_do_responsavel(client):
    """Cada um manda da própria fila: o convite tem que se apresentar com o nome
    de quem está mandando, não sempre "Rafael"."""
    fase = _fase_com(client, {"MG": 3, "ES": 3, "RJ": 3})
    _dividir(client, fase)
    for t in client.get("/api/tecnicos", params={"fase_id": fase["id"]}).json():
        for tipo in ("convite", "acionamento"):
            msg = client.get(f"/api/tecnicos/{t['id']}/mensagem", params={"tipo": tipo}).json()["mensagem"]
            assert f"Aqui é o {t['responsavel']}, da Faiston." in msg


def test_mensagem_sem_responsavel_continua_rafael(client):
    tec = client.post("/api/tecnicos", json={"nome": "Sem Dono", "telefone": "11999990000"}).json()
    msg = client.get(f"/api/tecnicos/{tec['id']}/mensagem").json()["mensagem"]
    assert "Aqui é o Rafael, da Faiston." in msg
    client.patch(f"/api/tecnicos/{tec['id']}", json={"responsavel": "Renato"})
    msg = client.get(f"/api/tecnicos/{tec['id']}/mensagem").json()["mensagem"]
    assert "Aqui é o Renato, da Faiston." in msg
