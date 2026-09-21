"""Os dois níveis de teste: a validação técnica e a da operação.

O que motivou isso: um caso passava na validação técnica, virava "Aprovado" e
sumia da pauta — mas na operação o mesmo caso não funcionava. Agora cada caso é
testado duas vezes, em níveis separados, e só fica Aprovado de verdade quando
os dois passam. Estes testes seguram exatamente essa promessa.
"""
import pytest

from app import niveis


# --------------------------------------------------------------- a regra

@pytest.mark.parametrize("interno,operacao,esperado", [
    ("Aprovado", "Aprovado", "Aprovado"),
    # o ponto central: passar na técnica não é passar na operação
    ("Aprovado", "Não testado", "Pendente na operação"),
    ("Não testado", "Aprovado", "Pendente na técnica"),
    ("Não testado", "Não testado", "Não testado"),
    # problema em qualquer nível derruba o caso inteiro
    ("Aprovado", "Reprovado", "Reprovado"),
    ("Reprovado", "Aprovado", "Reprovado"),
    ("Aprovado", "Bloqueado", "Bloqueado"),
    ("Bloqueado", "Reprovado", "Reprovado"),
    # N/A num nível = aquele nível não se aplica; quem decide é o outro
    ("N/A", "Aprovado", "Aprovado"),
    ("Aprovado", "N/A", "Aprovado"),
    ("N/A", "Não testado", "Não testado"),
    ("N/A", "N/A", "N/A"),
])
def test_consolidado_dos_dois_niveis(interno, operacao, esperado):
    assert niveis.status_geral(interno, operacao) == esperado


def test_nivel_aceita_os_apelidos_do_dia_a_dia():
    assert niveis.normaliza_nivel(None) == "interno"
    assert niveis.normaliza_nivel("operação") == "operacao"
    with pytest.raises(ValueError):
        niveis.normaliza_nivel("qualquer coisa")


# --------------------------------------------------------------- casos de teste

@pytest.fixture
def caso(client):
    """Um caso novo, criado pela tela — zerado nos dois níveis."""
    resp = client.post("/api/cases", json={
        "estagio": "01 · Níveis de teste",
        "resultado_esperado": "O caso guarda a validação técnica e a da operação separadas",
    })
    assert resp.status_code == 201
    yield resp.json()
    client.delete(f"/api/cases/{resp.json()['code']}")


def test_caso_nasce_zerado_nos_dois_niveis(caso):
    assert caso["status"] == "Não testado"
    assert caso["status_operacao"] == "Não testado"
    assert caso["status_geral"] == "Não testado"


def test_aprovado_so_na_tecnica_nao_aprova_o_caso(client, caso):
    code = caso["code"]
    resp = client.patch(f"/api/cases/{code}", json={"status": "Aprovado", "testado_por": "Rafael"})
    dados = resp.json()
    assert dados["status"] == "Aprovado"
    assert dados["status_operacao"] == "Não testado"
    # é isso que antes virava "Aprovado" e enganava a reunião
    assert dados["status_geral"] == "Pendente na operação"
    assert dados["testado_em"] is not None

    dados = client.patch(f"/api/cases/{code}", json={
        "status_operacao": "Aprovado", "testado_por_operacao": "Operação",
    }).json()
    assert dados["status_geral"] == "Aprovado"
    # cada nível guarda quem testou; um não sobrescreve o outro
    assert dados["testado_por"] == "Rafael"
    assert dados["testado_por_operacao"] == "Operação"
    assert dados["testado_em_operacao"] is not None


def test_reprovado_na_operacao_derruba_o_caso_aprovado_na_tecnica(client, caso):
    code = caso["code"]
    client.patch(f"/api/cases/{code}", json={"status": "Aprovado"})
    dados = client.patch(f"/api/cases/{code}", json={"status_operacao": "Reprovado"}).json()
    assert dados["status_geral"] == "Reprovado"
    # a validação técnica continua registrada como estava — o que falhou foi a operação
    assert dados["status"] == "Aprovado"


def test_chamado_de_cada_nivel_vive_separado(client, caso):
    code = caso["code"]
    dados = client.patch(f"/api/cases/{code}", json={
        "chamado": "221193", "chamado_operacao": "221207",
    }).json()
    assert dados["chamado"] == "221193"
    assert dados["chamado_operacao"] == "221207"


def test_status_invalido_no_nivel_operacao_e_recusado(client, caso):
    resp = client.patch(f"/api/cases/{caso['code']}", json={"status_operacao": "Quase"})
    assert resp.status_code == 400


def test_observacao_sabe_de_qual_teste_veio(client, caso):
    code = caso["code"]
    client.post(f"/api/cases/{code}/observacoes", json={"texto": "Funcionou no meu ambiente"})
    dados = client.post(f"/api/cases/{code}/observacoes", json={
        "texto": "Na operação travou na hora de salvar", "nivel": "operacao",
    }).json()
    assert [o["nivel"] for o in dados["observations"]] == ["interno", "operacao"]


def test_resumo_separa_o_consolidado_de_cada_nivel(client, caso):
    code = caso["code"]
    client.patch(f"/api/cases/{code}", json={"status": "Aprovado"})
    resumo = client.get("/api/summary").json()

    # o caso conta como "Pendente na operação" no consolidado, mas como aprovado
    # na leitura do nível técnico — as duas coisas são verdade ao mesmo tempo
    assert resumo["counts"]["Pendente na operação"] >= 1
    assert resumo["counts_interno"]["Aprovado"] >= 1
    assert resumo["counts_operacao"]["Não testado"] >= 1
    assert resumo["total"] == sum(resumo["counts"].values())
    assert resumo["pct_executado_operacao"] <= resumo["pct_executado_interno"]


def test_exportacao_mostra_os_dois_niveis_e_o_consolidado(client, caso):
    from io import BytesIO

    from openpyxl import load_workbook

    client.patch(f"/api/cases/{caso['code']}", json={"status": "Aprovado"})
    resp = client.get("/api/export?fluxo=C")
    assert resp.status_code == 200
    ws = load_workbook(BytesIO(resp.content)).active
    cabecalho = [c.value for c in ws[3]]
    assert cabecalho[-3:] == ["Validação técnica", "Validação na operação", "Status"]
    # a linha do caso traz o consolidado na última coluna, não o meu status
    linha = next(r for r in ws.iter_rows(values_only=True) if r[3] == caso["resultado_esperado"])
    assert linha[-3:] == ("Aprovado", "Não testado", "Pendente na operação")


# --------------------------------------------------------------- situações

@pytest.fixture
def estagio(client):
    """Uma situação com um estágio — o mesmo modelo de dois níveis."""
    sit = client.post("/api/situacoes", json={
        "titulo": "Chamado sem aceite", "descricao": "O técnico não aceita o chamado",
    }).json()
    sit = client.post(f"/api/situacoes/{sit['code']}/estagios", json={
        "nome": "01 · Criar chamado", "resultado_esperado": "Chamado criado",
    }).json()
    yield sit["code"], sit["estagios"][0]["id"]
    client.delete(f"/api/situacoes/{sit['code']}")


def test_estagio_de_situacao_tambem_tem_os_dois_niveis(client, estagio):
    code, estagio_id = estagio
    sit = client.patch(f"/api/situacoes/{code}/estagios/{estagio_id}", json={
        "status": "Aprovado", "testado_por": "Rafael",
    }).json()
    est = sit["estagios"][0]
    assert est["status_geral"] == "Pendente na operação"

    sit = client.patch(f"/api/situacoes/{code}/estagios/{estagio_id}", json={
        "status_operacao": "Aprovado", "testado_por_operacao": "Operação",
    }).json()
    est = sit["estagios"][0]
    assert est["status_geral"] == "Aprovado"
    assert est["testado_por_operacao"] == "Operação"


def test_situacao_guarda_o_chamado_da_operacao(client, estagio):
    code, _ = estagio
    sit = client.patch(f"/api/situacoes/{code}", json={
        "chamado": "221193", "chamado_operacao": "221207",
    }).json()
    assert (sit["chamado"], sit["chamado_operacao"]) == ("221193", "221207")


def test_pauta_cobra_o_que_so_passou_na_tecnica(client, caso):
    """O caso aprovado só na técnica tem que continuar aparecendo na reunião."""
    code = caso["code"]
    client.patch(f"/api/cases/{code}", json={"status": "Aprovado"})
    html = client.get("/relatorio").text
    assert code in html
    assert "Pendente na operação" in html or "esperando a operação" in html


# --------------------------------------------------------------- gestão de ativos

@pytest.fixture
def ajuste(client):
    """Um ajuste da Gestão de Ativos recém-levantado."""
    resp = client.post("/api/ativos/ajustes", json={
        "titulo": "Entrada não calcula o total",
        "atual": "Soma errado quando tem desconto",
        "esperado": "Somar certo",
        "tipo": "Bug",
        "versao": "v9",
    })
    assert resp.status_code == 201
    yield resp.json()
    client.delete(f"/api/ativos/ajustes/{resp.json()['id']}")


def test_ajuste_entregue_nao_fecha_so_com_a_validacao_tecnica(client, ajuste):
    """O mesmo ponto cego do Dispatcher: a LP entrega, a técnica confere e o
    item sumia da pauta sem a operação ter usado."""
    aid = ajuste["id"]
    assert ajuste["validacao_geral"] == "Não testado"
    assert ajuste["fechado"] is False

    client.patch(f"/api/ativos/ajustes/{aid}", json={"status": "entregue"})
    dados = client.patch(f"/api/ativos/ajustes/{aid}", json={
        "validacao": "Aprovado", "validado_por": "Rafael",
    }).json()
    assert dados["validacao_geral"] == "Pendente na operação"
    assert dados["fechado"] is False        # continua na pauta
    assert dados["validado_em"] is not None

    dados = client.patch(f"/api/ativos/ajustes/{aid}", json={
        "validacao_operacao": "Aprovado", "validado_por_operacao": "Estoque",
    }).json()
    assert dados["validacao_geral"] == "Aprovado"
    assert dados["fechado"] is True
    assert dados["validado_por"] == "Rafael"
    assert dados["validado_por_operacao"] == "Estoque"


def test_reprovado_na_operacao_reabre_o_ajuste(client, ajuste):
    aid = ajuste["id"]
    client.patch(f"/api/ativos/ajustes/{aid}", json={"validacao": "Aprovado"})
    dados = client.patch(f"/api/ativos/ajustes/{aid}", json={"validacao_operacao": "Reprovado"}).json()
    assert dados["validacao_geral"] == "Reprovado"
    assert dados["fechado"] is False


def test_status_validado_antigo_vira_entregue_mais_validacao_tecnica(client, ajuste):
    """Quem ainda manda o status "validado" (chamada antiga) está dizendo que a
    técnica aprovou — o ajuste não pode fechar por isso sozinho."""
    dados = client.patch(f"/api/ativos/ajustes/{ajuste['id']}", json={"status": "validado"}).json()
    assert dados["status"] == "entregue"
    assert dados["validacao"] == "Aprovado"
    assert dados["fechado"] is False


def test_validacao_invalida_e_recusada(client, ajuste):
    resp = client.patch(f"/api/ativos/ajustes/{ajuste['id']}", json={"validacao_operacao": "Mais ou menos"})
    assert resp.status_code == 400


def test_nota_do_ajuste_sabe_de_qual_validacao_veio(client, ajuste):
    aid = ajuste["id"]
    client.post(f"/api/ativos/ajustes/{aid}/observacoes", json={"texto": "Conferido na tela"})
    dados = client.post(f"/api/ativos/ajustes/{aid}/observacoes", json={
        "texto": "No estoque continua somando errado", "nivel": "operacao",
    }).json()
    assert [o["nivel"] for o in dados["observations"]] == ["interno", "operacao"]


def test_pauta_mantem_ajuste_que_so_a_tecnica_validou(client, ajuste):
    aid = ajuste["id"]
    client.patch(f"/api/ativos/ajustes/{aid}", json={"status": "entregue", "validacao": "Aprovado"})
    assert ajuste["titulo"] in client.get("/relatorio").text
